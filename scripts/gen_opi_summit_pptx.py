#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
# SPDX-FileCopyrightText: 2026 The dash-opi-map Authors
"""Build docs/opi-summit-2026-slides.pptx from the summit deck content."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "opi-summit-2026-slides.pptx"

INK = RGBColor(0x1B, 0x23, 0x2C)
MUTE = RGBColor(0x5C, 0x66, 0x72)
LINE = RGBColor(0xCF, 0xC8, 0xBA)
PAPER = RGBColor(0xF7, 0xF3, 0xEA)
CARD = RGBColor(0xFF, 0xFD, 0xF8)
TEAL = RGBColor(0x0C, 0x65, 0x6C)
TEAL_SOFT = RGBColor(0xD5, 0xE8, 0xE8)
NAVY = RGBColor(0x24, 0x3A, 0x56)
NAVY_SOFT = RGBColor(0xD9, 0xE2, 0xEE)
RUST_SOFT = RGBColor(0xF3, 0xDD, 0xD3)
CHIP = RGBColor(0xE7, 0xE1, 0xD4)
CREAM = RGBColor(0xF7, 0xF3, 0xEA)
NAVY_MUTE = RGBColor(0xC5, 0xD0, 0xDE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SANS = "Calibri"
MONO = "Consolas"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def rgb_hex(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def set_run(run, text, size, color, bold=False, font=SANS):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    run.font.italic = False


def _set_anchor(tf, anchor):
    tf.word_wrap = True
    body = tf._tt_txBody if hasattr(tf, "_tt_txBody") else tf._txBody
    bodyPr = body.find(qn("a:bodyPr"))
    if bodyPr is None:
        return
    mapping = {
        MSO_ANCHOR.TOP: "t",
        MSO_ANCHOR.MIDDLE: "ctr",
        MSO_ANCHOR.BOTTOM: "b",
    }
    bodyPr.set("anchor", mapping.get(anchor, "t"))


def add_rect(slide, l, t, w, h, fill, line=None, radius=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    return shape


def add_tb(
    slide,
    l,
    t,
    w,
    h,
    lines,
    *,
    size=14,
    color=INK,
    bold=False,
    font=SANS,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    fill=None,
    line=None,
):
    """lines: str or list of (text, size, color, bold, font)."""
    box = slide.shapes.add_textbox(l, t, w, h)
    if fill is not None:
        # textboxes cannot fill easily; use a rect behind instead
        pass
    tf = box.text_frame
    tf.word_wrap = True
    _set_anchor(tf, anchor)
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    if isinstance(lines, str):
        lines = [(lines, size, color, bold, font)]
    for i, item in enumerate(lines):
        if isinstance(item, str):
            item = (item, size, color, bold, font)
        text, sz, col, bld, fnt = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        set_run(run, text, sz, col, bld, fnt)
    return box


def cell(
    slide,
    l,
    t,
    w,
    h,
    title,
    sub=None,
    *,
    fill=CARD,
    line=LINE,
    title_color=INK,
    sub_color=MUTE,
    title_size=15,
    sub_size=11,
    mono_sub=False,
    center=False,
):
    add_rect(slide, l, t, w, h, fill, line)
    align = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    pad = Inches(0.12)
    lines = [(title, title_size, title_color, True, SANS)]
    if sub:
        lines.append((sub, sub_size, sub_color, False, MONO if mono_sub else SANS))
    add_tb(
        slide,
        l + pad,
        t + Inches(0.08),
        w - 2 * pad,
        h - Inches(0.16),
        lines,
        align=align,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def kicker_title(slide, kicker, title, n, nslides=12):
    add_tb(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.32),
           kicker.upper(), size=11, color=TEAL, bold=True)
    add_tb(slide, Inches(0.55), Inches(0.55), Inches(12.2), Inches(1.15),
           title, size=24, color=INK, bold=True)
    chrome(slide, n, nslides)


def chrome(slide, n, nslides=12):
    add_rect(slide, Inches(0), Inches(7.42), SLIDE_W, Inches(0.08), CHIP, line=None, radius=0)
    add_rect(slide, Inches(0), Inches(7.42), Emu(int(SLIDE_W * n / nslides)), Inches(0.08), TEAL, line=None, radius=0)
    add_tb(slide, Inches(0.55), Inches(7.12), Inches(8), Inches(0.28),
           "OPI Summit on DPU/IPUs 2026  ·  San Jose", size=10, color=MUTE)
    add_tb(slide, Inches(10.4), Inches(7.12), Inches(2.4), Inches(0.28),
           f"{n}  /  {nslides}", size=10, color=INK, bold=True, align=PP_ALIGN.RIGHT)


def pill(slide, l, t, text, w=None):
    w = w or Inches(0.16 * len(text) + 0.35)
    add_rect(slide, l, t, w, Inches(0.34), CHIP, line=None, radius=0.5)
    add_tb(slide, l, t, w, Inches(0.34), text, size=11, color=NAVY, bold=True,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def tag(slide, l, t, w, h, text):
    add_rect(slide, l, t, w, h, CARD, LINE, radius=0.12)
    add_tb(slide, l + Inches(0.04), t, w - Inches(0.08), h, text,
           size=10, color=INK, font=MONO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def chev(slide, l, t, h, text="→"):
    add_tb(slide, l, t, Inches(0.45), h, text, size=22, color=TEAL, bold=True,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def stack_cells(slide, l, t, w, h, items, gap=0.1):
    """items: list of kwargs for cell() minus l,t,w,h."""
    n = len(items)
    cell_h = (h - Inches(gap) * (n - 1)) / n
    y = t
    for item in items:
        cell(slide, l, y, w, cell_h, **item)
        y += cell_h + Inches(gap)


def panel(slide, l, t, w, h, heading):
    add_rect(slide, l, t, w, h, CARD, LINE, radius=0.08)
    add_tb(slide, l + Inches(0.14), t + Inches(0.1), w - Inches(0.28), Inches(0.32),
           heading, size=14, color=INK, bold=True)


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    return s


def add_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = text.strip()


def slide_1(prs):
    s = blank(prs)
    add_tb(s, Inches(0.55), Inches(0.42), Inches(12), Inches(0.32),
           "OPI SUMMIT ON DPU/IPUS  ·  15 OCTOBER 2026", size=12, color=TEAL, bold=True)
    add_tb(s, Inches(0.55), Inches(0.85), Inches(12), Inches(0.9),
           "Keep the SDK on the Card", size=40, color=INK, bold=True)
    add_tb(s, Inches(0.55), Inches(1.85), Inches(10), Inches(0.7),
           "SONiC DASH overlay objects, sent to the DPU over OPI CloudInfra gRPC.",
           size=20, color=INK)
    pill(s, Inches(0.55), Inches(2.65), "Standards and Interoperability", Inches(3.15))
    pill(s, Inches(3.85), Inches(2.65), "SONiC DASH × OPI", Inches(2.2))
    add_tb(s, Inches(0.55), Inches(3.2), Inches(8), Inches(0.7),
           [("Sridhar K N Rao", 16, MUTE, False, SANS),
            ("github.com/sknrao/dash-opi-map", 15, MUTE, False, MONO)])
    y, h = Inches(4.55), Inches(1.7)
    cell(s, Inches(0.55), y, Inches(3.55), h, "SONiC DASH", "SAI C ABI on the host",
         fill=NAVY_SOFT, line=RGBColor(0xC5, 0xD0, 0xDE), title_size=20, sub_size=13)
    chev(s, Inches(4.2), y, h)
    cell(s, Inches(4.7), y, Inches(3.55), h, "Adapter", "libsai_opi_adapter.so",
         fill=TEAL_SOFT, line=RGBColor(0xB7, 0xD4, 0xD4), title_size=20, sub_size=13, mono_sub=True)
    chev(s, Inches(8.35), y, h)
    cell(s, Inches(8.85), y, Inches(3.9), h, "OPI on the DPU", "CloudInfra + telemetry",
         fill=RUST_SOFT, line=RGBColor(0xE4, 0xC4, 0xB4), title_size=20, sub_size=13)
    chrome(s, 1)
    return s


def slide_2(prs):
    s = blank(prs)
    kicker_title(s, "The problem",
                 "DASH programs overlays through SAI on the host. DPU vendors program the card with their own SDK.", 2)
    add_tb(s, Inches(0.55), Inches(1.7), Inches(5.5), Inches(4.9), [
        ("SONiC DASH already has a C API on the host for tenant objects: ENI, VNET, routes, ACL. dashorch emits those as SAI calls.", 16, MUTE, False, SANS),
        ("", 8, MUTE, False, SANS),
        ("Those objects have to land in forwarding tables on the DPU. If you copy the switch-ASIC model, that pulls a vendor libsai.so and its SDK into the host syncd process.", 16, MUTE, False, SANS),
        ("", 8, MUTE, False, SANS),
        ("We do not want the vendor SDK on the host. This picture is the problem, not the design.", 16, MUTE, False, SANS),
    ])
    panel(s, Inches(6.35), Inches(1.7), Inches(3.05), Inches(4.95), "Host")
    stack_cells(s, Inches(6.52), Inches(2.15), Inches(2.71), Inches(4.3), [
        {"title": "dashorch", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE)},
        {"title": "syncd", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE)},
        {"title": "vendor libsai.so", "sub": "dlopen from the host", "fill": RUST_SOFT, "line": RGBColor(0xE4, 0xC4, 0xB4)},
        {"title": "proprietary SDK", "sub": "DOCA · IPDK · DSC · …", "fill": RUST_SOFT, "line": RGBColor(0xE4, 0xC4, 0xB4)},
    ])
    chev(s, Inches(9.45), Inches(1.7), Inches(4.95))
    panel(s, Inches(9.95), Inches(1.7), Inches(2.8), Inches(4.95), "DPU")
    stack_cells(s, Inches(10.12), Inches(2.15), Inches(2.46), Inches(4.3), [
        {"title": "SoC runtime", "sub": "drivers, firmware", "fill": TEAL_SOFT, "line": RGBColor(0xB7, 0xD4, 0xD4)},
        {"title": "ASIC tables", "sub": "overlay hardware", "fill": NAVY, "line": NAVY,
         "title_color": CREAM, "sub_color": NAVY_MUTE},
    ])
    return s


def slide_3(prs):
    s = blank(prs)
    kicker_title(s, "What has to land on the card",
                 "A tenant overlay is four kinds of object, written into DPU tables.", 3)
    # objects grid
    cell(s, Inches(2.55), Inches(1.8), Inches(4.7), Inches(1.05), "VNET",
         "tenant overlay · VXLAN VNI", fill=TEAL_SOFT, line=RGBColor(0xB7, 0xD4, 0xD4),
         title_size=18, center=True)
    cell(s, Inches(0.55), Inches(3.05), Inches(2.85), Inches(1.45), "ENI",
         "workload vNIC", fill=NAVY_SOFT, line=RGBColor(0xC5, 0xD0, 0xDE), title_size=18, center=True)
    cell(s, Inches(3.5), Inches(3.05), Inches(2.85), Inches(1.45), "Routes",
         "group + LPM entries", fill=NAVY_SOFT, line=RGBColor(0xC5, 0xD0, 0xDE), title_size=18, center=True)
    cell(s, Inches(6.45), Inches(3.05), Inches(2.85), Inches(1.45), "ACL",
         "group + rules", fill=NAVY_SOFT, line=RGBColor(0xC5, 0xD0, 0xDE), title_size=18, center=True)
    cell(s, Inches(0.55), Inches(4.7), Inches(8.75), Inches(1.55), "DPU ASIC overlay tables",
         fill=NAVY, line=NAVY, title_color=CREAM, title_size=20, center=True)
    add_rect(s, Inches(9.55), Inches(1.8), Inches(3.2), Inches(4.45), CARD, LINE)
    add_tb(s, Inches(9.72), Inches(2.05), Inches(2.86), Inches(4.0), [
        ("Where the objects come from: a cloud controller, or SAI-Challenger in the lab, writes DASH config into SONiC over gNMI. dashorch turns that into SAI calls.", 15, INK, False, SANS),
        ("", 10, MUTE, False, SANS),
        ("The missing hop is after SAI: a common way to land those objects on the DPU without a vendor SAI library on the host.", 15, MUTE, False, SANS),
    ])
    return s


def slide_4(prs):
    s = blank(prs)
    kicker_title(s, "Why this hurts",
                 "If DASH stays on host SAI only, each DPU vendor would still have to ship a different SAI library into SONiC.", 4)
    vendors = [
        ("Vendor A rack", "libsai A", "DPU A firmware"),
        ("Vendor B rack", "libsai B", "DPU B firmware"),
        ("Vendor C rack", "libsai C", "DPU C firmware"),
    ]
    x = Inches(0.55)
    w = Inches(3.9)
    for heading, lib, fw in vendors:
        panel(s, x, Inches(1.75), w, Inches(4.35), heading)
        stack_cells(s, x + Inches(0.16), Inches(2.2), w - Inches(0.32), Inches(3.7), [
            {"title": "SONiC + DASH", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE)},
            {"title": lib, "fill": RUST_SOFT, "line": RGBColor(0xE4, 0xC4, 0xB4)},
            {"title": fw, "fill": NAVY, "line": NAVY, "title_color": CREAM},
        ])
        x += w + Inches(0.22)
    add_tb(s, Inches(0.55), Inches(6.25), Inches(12.2), Inches(0.7),
           "DPU vendors do not do this today. They ship DOCA, IPDK, DSC on the card. This is what SONiC DASH would force if we treated the DPU like a switch ASIC.",
           size=15, color=MUTE)
    return s


def slide_5(prs):
    s = blank(prs)
    kicker_title(s, "How we solve it",
                 "Load an open SAI library on the host. Speak OPI to the card.", 5)
    panel(s, Inches(0.55), Inches(1.75), Inches(4.55), Inches(4.95), "Host")
    stack_cells(s, Inches(0.72), Inches(2.2), Inches(4.21), Inches(4.25), [
        {"title": "dashorch", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE)},
        {"title": "syncd", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE)},
        {"title": "libsai_opi_adapter.so", "sub": "loaded as libsai.so",
         "fill": TEAL_SOFT, "line": RGBColor(0xB7, 0xD4, 0xD4), "mono_sub": True},
    ])
    add_tb(s, Inches(5.15), Inches(1.75), Inches(1.15), Inches(4.95),
           [("gRPC", 16, TEAL, True, SANS), (":50051", 14, TEAL, False, MONO)],
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    panel(s, Inches(6.35), Inches(1.75), Inches(2.95), Inches(4.95), "DPU SoC")
    stack_cells(s, Inches(6.52), Inches(2.2), Inches(2.61), Inches(4.25), [
        {"title": "OPI bridge", "sub": "CloudInfraService",
         "fill": TEAL_SOFT, "line": RGBColor(0xB7, 0xD4, 0xD4), "mono_sub": True},
        {"title": "vendor SDK", "sub": "on the card", "fill": RUST_SOFT, "line": RGBColor(0xE4, 0xC4, 0xB4)},
        {"title": "DPU ASIC", "fill": NAVY, "line": NAVY, "title_color": CREAM},
    ])
    notes = [
        ("Same load path SONiC already has.",
         "syncd still calls dlopen(\"libsai.so\"). The file behind that name is the adapter."),
        ("One hop off the host.",
         "SAI calls are packed into OPI protobuf and sent to the DPU ARM or x86 complex."),
        ("Vendor runtime stays on the SoC.",
         "DOCA, IPDK, and the rest run behind the OPI server."),
    ]
    y = Inches(1.75)
    for title, body in notes:
        add_rect(s, Inches(9.5), y, Inches(3.28), Inches(1.55), CARD, LINE)
        add_tb(s, Inches(9.64), y + Inches(0.12), Inches(3.0), Inches(1.35), [
            (title, 13, INK, True, SANS),
            (body, 12, MUTE, False, SANS),
        ])
        y += Inches(1.67)
    return s


def slide_6(prs):
    s = blank(prs)
    kicker_title(s, "Architecture", "Four tiers, from tenant intent to ASIC tables.", 6)
    cols = [
        ("1 · Intent", [
            {"title": "Cloud SDN", "sub": "or SAI-Challenger", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE)},
            {"title": "gNMI Set", "sub": "DASH YANG / JSON", "mono_sub": True},
        ]),
        ("2 · SONiC", [
            {"title": "sonic-gnmi", "sub": ":8080", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE), "mono_sub": True},
            {"title": "APPL_DB", "sub": "DASH_*_TABLE", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE), "mono_sub": True},
            {"title": "dashorch", "sub": "SAI ops", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE)},
        ]),
        ("3 · Adapter", [
            {"title": "syncd", "fill": NAVY_SOFT, "line": RGBColor(0xC5, 0xD0, 0xDE)},
            {"title": "adapter", "sub": "libsai.so · same process", "fill": TEAL_SOFT, "line": RGBColor(0xB7, 0xD4, 0xD4), "mono_sub": True},
            {"title": "sai_create_eni", "sub": "sai_get_eni_stats", "mono_sub": True},
        ]),
        ("4 · DPU", [
            {"title": "OPI gRPC", "sub": ":50051", "fill": TEAL_SOFT, "line": RGBColor(0xB7, 0xD4, 0xD4), "mono_sub": True},
            {"title": "vendor bridge", "sub": "on the SoC", "fill": RUST_SOFT, "line": RGBColor(0xE4, 0xC4, 0xB4)},
            {"title": "ASIC", "fill": NAVY, "line": NAVY, "title_color": CREAM},
        ]),
    ]
    x = Inches(0.45)
    w = Inches(3.0)
    for i, (heading, items) in enumerate(cols):
        line = TEAL if i == 2 else LINE
        add_rect(s, x, Inches(1.75), w, Inches(5.0), CARD, line, radius=0.08)
        add_tb(s, x + Inches(0.14), Inches(1.85), w - Inches(0.28), Inches(0.35),
               heading, size=14, color=INK, bold=True)
        stack_cells(s, x + Inches(0.14), Inches(2.3), w - Inches(0.28), Inches(4.2), items)
        x += w + Inches(0.16)
    return s


def slide_7(prs):
    s = blank(prs)
    kicker_title(s, "OPI tie-in",
                 "The adapter does not call DOCA or IPDK. It calls the published OPI CloudInfra API.", 7)
    # 2x2
    boxes = [
        (Inches(0.5), Inches(1.7), Inches(7.15), Inches(2.55), TEAL_SOFT, RGBColor(0xB7, 0xD4, 0xD4),
         "CloudInfraService",
         "opi_api.network.cloud.v1alpha1  ·  listener :50051",
         ["CreateVnic", "CreateInterface", "CreateVpc", "CreateSubnet",
          "CreateRouteTable", "CreateRoute", "CreateSecurityPolicy",
          "CreateSecurityRule", "CreateMapping", "Get / Update / Delete / List"]),
        (Inches(7.85), Inches(1.7), Inches(4.95), Inches(2.55), NAVY_SOFT, RGBColor(0xC5, 0xD0, 0xDE),
         "Resources we program",
         "Cloud overlay model on the DPU.",
         ["Vnic + VnicSpec", "Interface / HostIfSpec", "Vpc + Subnet",
          "RouteTable + Route", "SecurityPolicy + SecurityRule", "Mapping"]),
        (Inches(0.5), Inches(4.4), Inches(7.15), Inches(2.4), RUST_SOFT, RGBColor(0xE4, 0xC4, 0xB4),
         "TelemetryService",
         "Same gRPC server as CloudInfra. FlexCounter still calls sai_get_eni_stats.",
         ["GetVnicMetrics", "rx_bytes / tx_packets", "policy_drops", "forwarding_drops"]),
        (Inches(7.85), Inches(4.4), Inches(4.95), Inches(2.4), CARD, LINE,
         "Shared types",
         "opi_api.network.opinetcommon.v1alpha1",
         ["Encap / VXLAN vnid", "IPPrefix", "RuleMatch", "SecurityRuleAction"]),
    ]
    for l, t, w, h, fill, line, title, sub, tags in boxes:
        add_rect(s, l, t, w, h, fill, line)
        add_tb(s, l + Inches(0.16), t + Inches(0.1), w - Inches(0.32), Inches(0.32),
               title, size=16, color=INK, bold=True, font=MONO if "Service" in title or title == "Shared types" else SANS)
        add_tb(s, l + Inches(0.16), t + Inches(0.42), w - Inches(0.32), Inches(0.45),
               sub, size=12, color=MUTE, font=MONO if "opi_api" in sub or "sai_" in sub else SANS)
        tw, th = Inches(1.55), Inches(0.32)
        x, y = l + Inches(0.14), t + Inches(0.95)
        for tg in tags:
            need = Inches(max(1.35, 0.09 * len(tg) + 0.28))
            if x + need > l + w - Inches(0.12):
                x = l + Inches(0.14)
                y += Inches(0.4)
            tag(s, x, y, need, th, tg)
            x += need + Inches(0.08)
    return s


def slide_8(prs):
    s = blank(prs)
    kicker_title(s, "The contract",
                 "ENI, VNET, route, and ACL each have a matching OPI object. That mapping is this project.", 8)
    add_tb(s, Inches(0.55), Inches(1.7), Inches(5.6), Inches(0.3),
           "DASH SAI", size=12, color=MUTE, bold=True)
    add_tb(s, Inches(7.15), Inches(1.7), Inches(5.6), Inches(0.3),
           "OPI CLOUD", size=12, color=MUTE, bold=True)
    rows = [
        ("ENI", "workload interface · MAC, VNI, CPS",
         "Vnic + Interface", "vnic.proto · interface.proto"),
        ("VNET", "tenant overlay · 24-bit VNI",
         "Vpc + Subnet", "vpc.proto · subnet.proto · VPC_TYPE_TENANT"),
        ("Outbound routing", "group + LPM entry",
         "RouteTable + Route", "route.proto · RouteClass"),
        ("ACL", "group + match, action, priority",
         "SecurityPolicy + SecurityRule", "networkpolicy.proto"),
    ]
    y = Inches(2.05)
    h = Inches(1.12)
    for left_t, left_s, right_t, right_s in rows:
        cell(s, Inches(0.55), y, Inches(5.35), h, left_t, left_s,
             fill=NAVY_SOFT, line=RGBColor(0xC5, 0xD0, 0xDE), title_size=16)
        chev(s, Inches(6.05), y, h)
        cell(s, Inches(6.55), y, Inches(6.2), h, right_t, right_s,
             fill=TEAL_SOFT, line=RGBColor(0xB7, 0xD4, 0xD4), title_size=16, mono_sub=True)
        y += h + Inches(0.12)
    return s


def slide_9(prs):
    s = blank(prs)
    kicker_title(s, "A worked path",
                 "Creating an ENI means two gRPC calls to the DPU. The DPU then programs hardware.", 9)
    steps = [
        ("01", "gNMI Set", "DASH_ENI_TABLE into sonic-gnmi.", CARD, LINE, False),
        ("02", "dashorch", "sai_create_eni", CARD, LINE, True),
        ("03", "CreateInterface", "HostIfSpec for the PF/VF.", TEAL_SOFT, RGBColor(0xB7, 0xD4, 0xD4), True),
        ("04", "CreateVnic", "subnet, MAC, VXLAN VNI, max_cps.", TEAL_SOFT, RGBColor(0xB7, 0xD4, 0xD4), False),
        ("05", "Bridge → ASIC", "Vendor runtime writes tables.", RUST_SOFT, RGBColor(0xE4, 0xC4, 0xB4), False),
    ]
    x = Inches(0.4)
    w = Inches(2.15)
    h = Inches(3.55)
    y = Inches(1.75)
    for i, (num, title, body, fill, line, mono) in enumerate(steps):
        add_rect(s, x, y, w, h, fill, line)
        add_tb(s, x + Inches(0.12), y + Inches(0.18), w - Inches(0.24), Inches(0.4),
               num, size=14, color=TEAL, bold=True, font=MONO)
        add_tb(s, x + Inches(0.12), y + Inches(0.6), w - Inches(0.24), Inches(0.7),
               title, size=16, color=INK, bold=True)
        add_tb(s, x + Inches(0.12), y + Inches(1.35), w - Inches(0.24), Inches(1.8),
               body, size=13, color=MUTE, font=MONO if mono else SANS)
        if i < len(steps) - 1:
            chev(s, x + w, y, h, "›")
        x += w + Inches(0.32)
    add_tb(s, Inches(0.55), Inches(5.5), Inches(12.2), Inches(1.35),
           "Order on the wire follows OPI: Vpc → Subnet → Interface → Vnic. DASH binds the ENI to a VNET; the adapter walks VNET → subnet → VPC so VnicSpec.subnet_name_ref is valid.",
           size=15, color=MUTE)
    return s


def add_table(slide, l, t, w, h, headers, rows, col_w=None):
    table_shape = slide.shapes.add_table(len(rows) + 1, 2, l, t, w, h)
    table = table_shape.table
    if col_w:
        table.columns[0].width = col_w[0]
        table.columns[1].width = col_w[1]
    for i, hdr in enumerate(headers):
        cell_ = table.cell(0, i)
        cell_.text = ""
        p = cell_.text_frame.paragraphs[0]
        run = p.add_run()
        set_run(run, hdr, 11, MUTE, True, SANS)
        cell_.fill.solid()
        cell_.fill.fore_color.rgb = CARD
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell_ = table.cell(r, c)
            cell_.text = ""
            p = cell_.text_frame.paragraphs[0]
            run = p.add_run()
            set_run(run, val, 11, INK, False, MONO)
            cell_.fill.solid()
            cell_.fill.fore_color.rgb = CARD
    # borders
    tbl = table._tbl
    for tc in tbl.iter(qn("a:tcPr")):
        for edge in ("lnL", "lnR", "lnT", "lnB"):
            ln = etree.SubElement(tc, qn(f"a:{edge}"))
            ln.set("w", "6350")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            sc = etree.SubElement(sf, qn("a:srgbClr"))
            sc.set("val", "CFC8BA")
    return table_shape


def slide_10(prs):
    s = blank(prs)
    kicker_title(s, "Translation, in the small",
                 "A few SAI attributes, and the OPI fields they become.", 10)
    add_rect(s, Inches(0.5), Inches(1.7), Inches(6.2), Inches(5.05), CARD, LINE)
    add_table(
        s, Inches(0.62), Inches(1.85), Inches(5.96), Inches(4.75),
        ["DASH SAI", "OPI"],
        [
            ("SAI_ENI_ATTR_VNET_ID", "VnicSpec.subnet_name_ref → Subnet → Vpc"),
            ("SAI_ENI_ATTR_VM_VNI", "VnicSpec.fabric_encap · ENCAP_TYPE_VXLAN"),
            ("eni MAC", "VnicSpec.mac_address"),
            ("SAI_ENI_ATTR_CPS", "VnicSpec.max_cps"),
            ("SAI_ENI_ATTR_FLOWS", "VnicSpec.max_sessions"),
            ("ACL group ids", "ingress_* / egress_*_security_policy_name_ref"),
            ("routing group id", "v4/v6_route_table_name_ref"),
        ],
        col_w=(Inches(2.35), Inches(3.61)),
    )
    add_rect(s, Inches(6.9), Inches(1.7), Inches(5.9), Inches(5.05), CARD, LINE)
    add_table(
        s, Inches(7.02), Inches(1.85), Inches(5.66), Inches(2.85),
        ["sai_get_eni_stats", "GetVnicMetrics"],
        [
            ("SAI_ENI_STAT_RX_BYTES", "rx_bytes"),
            ("SAI_ENI_STAT_TX_PACKETS", "tx_packets"),
            ("FORWARDING_DROP_PACKETS", "forwarding_drops"),
            ("POLICY_DROP_PACKETS", "policy_drops"),
        ],
        col_w=(Inches(3.0), Inches(2.66)),
    )
    add_tb(s, Inches(7.1), Inches(4.9), Inches(5.5), Inches(1.5),
           "OPI inventory.proto is SMBIOS and FRU. Packet counters are TelemetryService on the same CloudInfra port.",
           size=14, color=MUTE)
    return s


def slide_11(prs):
    s = blank(prs)
    kicker_title(s, "Why this is worth doing here", "One mapping, three audiences.", 11)
    cols = [
        (TEAL_SOFT, RGBColor(0xB7, 0xD4, 0xD4), "OPI",
         "Cloud Vnic, Vpc, Route, and SecurityPolicy carrying a real NOS overlay. Dataplane metrics sit on TelemetryService, next to CloudInfra."),
        (NAVY_SOFT, RGBColor(0xC5, 0xD0, 0xDE), "SONiC DASH",
         "The host ABI stays SAI. dashorch does not link a vendor SDK. OPI is a SAI backend the NOS can load the way it loads any other libsai.so."),
        (RUST_SOFT, RGBColor(0xE4, 0xC4, 0xB4), "Vendors and operators",
         "Ship CloudInfraService plus telemetry on the SoC. Keep the closed SDK there. The same SONiC image can follow any card that speaks that API."),
    ]
    x = Inches(0.5)
    w = Inches(4.0)
    for fill, line, who, body in cols:
        add_rect(s, x, Inches(1.75), w, Inches(5.0), fill, line)
        add_tb(s, x + Inches(0.2), Inches(1.95), w - Inches(0.4), Inches(0.4),
               who.upper(), size=13, color=TEAL, bold=True)
        add_tb(s, x + Inches(0.2), Inches(2.5), w - Inches(0.4), Inches(3.9),
               body, size=16, color=INK)
        x += w + Inches(0.16)
    return s


def slide_12(prs):
    s = blank(prs)
    kicker_title(s, "Asks", "What this room can do with it.", 12)
    asks = [
        ("On the card",
         "Implement CloudInfraService and TelemetryService.GetVnicMetrics. That is the DPU-side contract."),
        ("On the host",
         "Treat OPI as a SAI vendor library for DASH. Review the gRPC to SAI status map and FlexCounter behavior."),
        ("Together",
         "Stand up one lab with SONiC and two OPI bridges. The mapping lives in the repo so both sides can review the same object names."),
    ]
    y = Inches(1.75)
    for title, body in asks:
        add_rect(s, Inches(0.5), y, Inches(6.55), Inches(1.55), CARD, LINE)
        add_tb(s, Inches(0.68), y + Inches(0.14), Inches(6.2), Inches(1.28), [
            (title, 16, INK, True, SANS),
            (body, 14, MUTE, False, SANS),
        ])
        y += Inches(1.7)
    add_rect(s, Inches(7.25), Inches(1.75), Inches(5.55), Inches(5.0), NAVY, NAVY)
    add_tb(s, Inches(7.5), Inches(2.15), Inches(5.1), Inches(0.35),
           "CODE AND CONTRACT", size=12, color=RGBColor(0x9F, 0xD0, 0xD3), bold=True)
    add_tb(s, Inches(7.5), Inches(2.6), Inches(5.1), Inches(0.7),
           "github.com/sknrao/dash-opi-map", size=18, color=CREAM, font=MONO)
    add_tb(s, Inches(7.5), Inches(3.4), Inches(5.1), Inches(0.7),
           "Apache-2.0  ·  mapping YAML  ·  adapter  ·  compose lab",
           size=14, color=NAVY_MUTE)
    add_tb(s, Inches(7.5), Inches(4.7), Inches(5.1), Inches(0.8),
           "Questions", size=32, color=CREAM, bold=True)
    return s


NOTES = {
    1: """
One sentence: SONiC DASH already knows how to program tenant overlays. We load an open library on the host so those calls go to the DPU over OPI, and the vendor SDK never leaves the card.

What "OPI Cloud" means: it is not "the cloud." It is the CloudInfra API from the Linux Foundation OPI project. The objects are Vnic, Vpc, Subnet, Route, SecurityPolicy. They are defined in github.com/opiproject/opi-api, under network/cloud. When the slide says CloudInfra, that is this API, listening on port 50051 on the DPU.

The three boxes: SONiC still calls SAI on the host. The file behind libsai.so is our adapter, not NVIDIA or Intel. The adapter talks gRPC to an OPI server that runs on the DPU ARM or x86.
""",
    2: """
Do not present this as current DPU vendor practice. NVIDIA, Intel, and AMD are not shipping a SAI library into the SONiC image today. They already put DOCA, IPDK, or DSC on the card. This slide is the switch-ASIC pattern applied to a DPU: if DASH's only host interface is SAI, the vendor SDK gets pulled into the host process.

"Host ABI" = the C functions SONiC already calls on the x86 host: sai_create_eni, sai_create_vnet, and so on. We are not inventing a new host API. dashorch keeps calling SAI.

Why the proprietary SDK is shown on the host: that is the problem picture. We do not need it there, and we do not want it there. On a switch, the vendor SAI library on the host talks the vendor SDK. Copy that to a DPU and SONiC now links DOCA or IPDK. Slide 5 is the fix: SDK stays on the card.

"ASIC tables / overlay hardware" = the forwarding tables in the DPU silicon. VXLAN encap/decap, tenant route lookup, ACL. Not a software switch on the host, and not the ARM CPU. The packets have to be forwarded by that hardware.
""",
    3: """
Four objects, say them once: VNET is the tenant overlay (VXLAN VNI). ENI is the workload vNIC (VM or container). Routes are LPM entries in a group. ACL is match/action rules in a group. Those four have to be written into DPU hardware tables or the overlay does not exist on the wire.

"DPU ASIC overlay tables" is the same idea as "ASIC tables / overlay hardware" on the previous slide. It is the silicon forwarding tables, not Linux on the ARM.

Why the two lines on the right: they are only the northbound reminder. A cloud SDN controller — or SAI-Challenger, which is our lab test client — writes DASH YANG into SONiC over gNMI. dashorch reads APPL_DB and makes SAI calls. That path already exists. The gap we are here to close is the hop after SAI: get those objects onto the DPU without each vendor writing a host SAI plugin.
""",
    4: """
Say this out loud so you do not overclaim: DPU vendors are not doing this today. They are not dropping a different libsai.so into SONiC. They ship their SDK on the card.

The title is a warning, not a status report. SONiC DASH's host interface is SAI. Switch vendors already ship one SAI library per ASIC. If we do nothing, the only way to run DASH on a DPU is each vendor writing their own SAI plugin for SONiC. Then swapping a BlueField for an Intel IPU changes the NOS image, the package, and the failure domain. Mixed DPU SKUs become a SONiC problem.

OPI is how we avoid that: one host library, many cards, as long as the card speaks CloudInfra.
""",
    5: """
Vendor SDK and vendor / OPI bridge are not the same thing.

Vendor SDK on the card = NVIDIA DOCA, Intel IPDK, AMD Pensando DSC. The vendor's own runtime that programs the silicon. It belongs on the DPU ARM or x86. The host must not link it.

OPI bridge on the SoC = a gRPC server we (or the vendor) run on that same ARM. It implements CloudInfraService. In this repo that is dpu_cloudinfra_bridge on NVIDIA and the Intel Phase A daemon. The bridge receives CreateVnic and friends, then calls the vendor SDK. The host only sees OPI.

The three cards on the right: same dlopen path SONiC already has; the only network hop is gRPC :50051; DOCA/IPDK stay behind the OPI server.
""",
    6: """
Walk left to right in 30 seconds.

1. Intent: cloud controller or SAI-Challenger writes DASH config.
2. SONiC: gNMI into APPL_DB, dashorch emits SAI.
3. Adapter: still inside syncd, loaded as libsai.so. This is our code.
4. DPU: OPI gRPC to port 50051. The vendor bridge is the OPI server on the ARM. The vendor SDK sits behind that server. ASIC is the hardware tables.

If someone asks "is the bridge the SDK?": no. Bridge speaks OPI. SDK speaks the vendor. ASIC is silicon.
""",
    7: """
Say the title in plain words: the adapter on the host does not call NVIDIA or Intel APIs. It calls CloudInfraService, which is a published gRPC API in the OPI opi-api repository.

"Southbound" just means "toward the hardware." SAI stays the host-facing API. OPI CloudInfra is what we send to the DPU.

Shared types: yes, they exist today. They are not a service we wrote. They are reusable protobuf messages in opi-api, package opinetcommon (networktypes.proto): Encap (VXLAN VNI), IPPrefix, RuleMatch, SecurityRuleAction. CreateVnic, CreateRoute, and CreateSecurityRule fill those fields. There is no extra process named "shared types."
""",
    8: """
Say it without the metaphor: DASH and OPI use different object names for the same tenant overlay. This repository publishes that mapping and implements it.

ENI → Vnic plus Interface (the PF/VF).
VNET → Vpc plus Subnet.
Outbound routing group/entry → RouteTable plus Route.
ACL group/rule → SecurityPolicy plus SecurityRule.

That YAML and the adapter are the project. Vendors implement CloudInfra using these names, not DASH SAI names.
""",
    9: """
Walk the five boxes. When SONiC creates an ENI, the adapter does not write hardware itself. It makes two gRPC calls to the DPU: CreateInterface for the PF or VF, then CreateVnic for MAC, VNI, and CPS. The daemon on the DPU then programs forwarding tables. "Silicon" on the old title just meant those hardware tables.

Mention the order if you have 10 seconds: OPI wants Vpc, then Subnet, then Interface, then Vnic. DASH only says "this ENI belongs to this VNET." The adapter walks that chain so the Vnic points at a real subnet.
""",
    10: """
Do not read the table. Pick two rows: VM VNI becomes VXLAN encap on the Vnic, and ACL group ids become security-policy name refs.

Stats: FlexCounter still calls sai_get_eni_stats. We do not invent a second host API. Packet counters come from TelemetryService.GetVnicMetrics on the same CloudInfra port 50051. OPI Inventory is FRU/SMBIOS only. It is not dataplane stats.
""",
    11: """
OPI room: CloudInfra is carrying a real NOS overlay, not a demo proto.
SONiC DASH room: dashorch does not link DOCA. It still dlopens libsai.so.
Vendor room: implement CloudInfra plus GetVnicMetrics on the ARM. Keep DOCA/IPDK there. One SONiC image can follow any card that speaks that API.
""",
    12: """
Three asks, then repo, then stop.

On the card: implement CloudInfraService and GetVnicMetrics.
On the host: treat this adapter as a SAI vendor library for DASH. Review status mapping and FlexCounter.
Together: one lab, SONiC plus two vendor bridges, same object names.

github.com/sknrao/dash-opi-map
""",
}


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for i, fn in enumerate(
        (slide_1, slide_2, slide_3, slide_4, slide_5, slide_6,
         slide_7, slide_8, slide_9, slide_10, slide_11, slide_12),
        start=1,
    ):
        add_notes(fn(prs), NOTES[i])
    prs.save(OUT)
    print(f"wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
